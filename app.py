# ------------------------------------------------------------------------------
# Imports & Setup
# ------------------------------------------------------------------------------
import base64
import html
import sqlite3
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from sqlalchemy import text
from supabase import create_client

# ------------------------------------------------------------------------------
# Supabase & Database Initialization
# ------------------------------------------------------------------------------
supabase_settings = st.secrets["supabase"]
supabase_url = supabase_settings.get("url")
supabase_key = supabase_settings.get("key")
supabase_anon_key = supabase_settings.get("anon_key", supabase_key)
supabase_secret_key = supabase_settings.get("secret_key")

missing_supabase_secrets = []
if not supabase_url:
    missing_supabase_secrets.append("url")
if not supabase_anon_key:
    missing_supabase_secrets.append("anon_key or key")
if missing_supabase_secrets:
    st.error(
        "Missing [supabase] value(s) in .streamlit/secrets.toml: "
        + ", ".join(missing_supabase_secrets)
    )
    st.stop()

supabase = create_client(supabase_url, supabase_anon_key)
supabase_admin = (
    create_client(supabase_url, supabase_secret_key)
    if supabase_secret_key
    else None
)


def get_db_connection():
    """Returns Streamlit SQL Connection wrapper."""
    return st.connection("postgres", type="sql")


def init_db():
    """Initializes schema in Supabase PostgreSQL."""
    db_conn = get_db_connection()
    with db_conn.session as session:
        session.execute(
            text(
                """
            CREATE TABLE IF NOT EXISTS incidents (
                id SERIAL PRIMARY KEY,
                dashboard_name VARCHAR(100) NOT NULL,
                incident_date TIMESTAMP,
                barangay VARCHAR(100),
                service VARCHAR(100),
                traffic VARCHAR(50),
                registrant VARCHAR(100),
                photo_url TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
        """
            )
        )
        session.commit()


# Initialize PostgreSQL schema upon startup
try:
    init_db()
except Exception as err:
    st.warning(
        "The Incident database connection is not configured yet. Copy the "
        "[connections.postgres] section from your original Incident Dashboard "
        "secrets.toml into this project's .streamlit/secrets.toml."
    )


@st.cache_data(ttl=30)
def load_dashboard_data(dashboard_name=None):
    """Loads incidents from Supabase SQL table."""
    db_conn = get_db_connection()
    df = pd.DataFrame()

    try:
        if dashboard_name and dashboard_name != "Overall View":
            query = "SELECT * FROM incidents WHERE dashboard_name = :dname ORDER BY incident_date DESC;"
            df = db_conn.query(query, params={"dname": dashboard_name})
        else:
            df = db_conn.query("SELECT * FROM incidents ORDER BY incident_date DESC;")
    except Exception as e:
        print("Database query failed:", e)

    return df


def fetch_uploaded_files_from_supabase() -> dict[str, str]:
    """Fetches uploaded files from Supabase Storage bucket and maps dashboard names to filenames."""
    bucket = "dashboards"
    try:
        files = supabase.storage.from_(bucket).list()
        if not isinstance(files, list):
            return {}

        dashboard_files = {}
        for file_info in files:
            file_name = file_info.get("name", "")
            if "_" in file_name:
                # Extract the dashboard prefix before the first underscore
                dashboard_prefix = file_name.split("_")[0]
                dashboard_files[dashboard_prefix] = file_name
            else:
                dashboard_files[file_name] = file_name

        return dashboard_files
    except Exception as e:
        print(f"Error fetching files from Supabase storage: {e}")
        return {}


def download_dashboard_file_from_supabase(
    dashboard_name: str,
) -> tuple[bytes | None, str | None]:
    """Download the newest original CSV/Excel file saved for a dashboard."""
    bucket = "dashboards"
    expected_prefix = f"{dashboard_name}_".casefold()
    try:
        files = supabase.storage.from_(bucket).list()
        matches = [
            file_info
            for file_info in files
            if str(file_info.get("name", "")).casefold().startswith(expected_prefix)
            and Path(str(file_info.get("name", ""))).suffix.casefold()
            in {".csv", ".xlsx", ".xls"}
        ]
        if not matches:
            return None, None

        # Supabase returns created_at/updated_at metadata. Sorting also gives a
        # deterministic filename fallback when that metadata is unavailable.
        newest = max(
            matches,
            key=lambda item: (
                str(item.get("updated_at") or item.get("created_at") or ""),
                str(item.get("name", "")),
            ),
        )
        stored_name = str(newest["name"])
        file_bytes = supabase.storage.from_(bucket).download(stored_name)
        return bytes(file_bytes), stored_name
    except Exception as error:
        print(f"Stored dashboard download warning: {error}")
        return None, None


def upload_photo_to_supabase(file_bytes: bytes, filename: str) -> str:
    """Uploads photo or data file to Supabase Storage bucket and returns public URL."""
    bucket = "dashboards"
    path = f"{filename}"

    try:
        response = supabase.storage.from_(bucket).upload(
            path=path,
            file=file_bytes,
            file_options={"upsert": "true"}
        )
        if hasattr(response, "error") and response.error:
            print("Supabase Upload Error:", response.error)
            return ""

        return supabase.storage.from_(bucket).get_public_url(path)
    except Exception as e:
        print("Supabase Upload Exception:", e)
        return ""


def save_dataframe_to_supabase(df: pd.DataFrame, dashboard_name: str) -> None:
    """Appends records directly to Supabase Database."""
    db_conn = get_db_connection()
    df_to_save = df.copy()

    schema_mapping = {
        "Registered Date": "incident_date",
        "Barangay": "barangay",
        "Type of Incident": "service",
        "Traffic": "traffic",
        "Registrant": "registrant",
        "Photo": "photo_url",
    }
    df_to_save = df_to_save.rename(columns=schema_mapping)
    df_to_save["dashboard_name"] = dashboard_name

    allowed_cols = [
        "dashboard_name",
        "incident_date",
        "barangay",
        "service",
        "traffic",
        "registrant",
        "photo_url",
    ]
    valid_cols = [col for col in allowed_cols if col in df_to_save.columns]

    if valid_cols:
        engine = db_conn.engine
        df_to_save[valid_cols].to_sql(
            "incidents", con=engine, if_exists="append", index=False
        )
        load_dashboard_data.clear()


def delete_dashboard_data_from_supabase(
    dashboard_name: str, requested_by: str
) -> bool:
    """Delete a dashboard only when the signed-in account is an Admin or Super Admin."""
    # Check the current database role instead of trusting only the role cached in
    # Streamlit session state. This also applies role changes immediately.
    if not st.session_state.get("admin_authenticated"):
        st.error("You must be signed in to delete dashboard data.")
        return False

    signed_in_user = st.session_state.get("admin_username")
    if not signed_in_user or signed_in_user.casefold() != requested_by.casefold():
        st.error("The delete request could not be verified.")
        return False

    current_role = get_admin_role(signed_in_user)
    if current_role not in {"admin", "super_admin"}:
        st.error("Operators are not allowed to delete dashboard data.")
        return False

    bucket = "dashboards"
    db_conn = get_db_connection()
    try:
        # 1. Clean up records in Supabase PostgreSQL
        with db_conn.session as session:
            session.execute(
                text("DELETE FROM incidents WHERE dashboard_name = :dname;"),
                {"dname": dashboard_name},
            )
            session.commit()

        # 2. Delete matched files from Supabase Storage
        try:
            files = supabase.storage.from_(bucket).list()
            matched_files = [
                f["name"] for f in files
                if dashboard_name.casefold() in f["name"].casefold()
            ]
            if matched_files:
                supabase.storage.from_(bucket).remove(matched_files)
        except Exception as storage_err:
            print(f"Storage deletion warning: {storage_err}")

        # 3. Purge session state & clear cached database queries
        if dashboard_name in st.session_state.uploaded_dashboards:
            del st.session_state.uploaded_dashboards[dashboard_name]

        # FIX: Set pending_dashboard instead of directly setting selected_dashboard
        if st.session_state.get("selected_dashboard") == dashboard_name:
            st.session_state.pending_dashboard = OVERALL_DASHBOARD_NAME

        load_dashboard_data.clear()
        st.cache_data.clear()
        return True

    except Exception as e:
        st.error(f"Failed to delete dashboard records from database: {e}")
        return False


# ------------------------------------------------------------------------------
# Application Configuration & Constants
# ------------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).parent
DEFAULT_TITLE = "SOCOCA Dashboard"
OVERALL_DASHBOARD_NAME = "Overall View"
ADMIN_DB_FILE = Path(__file__).with_name("dashboard_admin.db")

st.set_page_config(page_title=DEFAULT_TITLE, layout="wide")

# CSS injection for dynamic theme adaptability
st.markdown(
    """
    <style>
    html, body, [class*="css"], [data-testid="stAppViewContainer"] {
        font-family: "MS PGothic", "ＭＳ Ｐゴシック", "MS Gothic",
                     "Yu Gothic", "Noto Sans JP", Arial, sans-serif;
    }
    .block-container {padding-top: 1.5rem; padding-bottom: 2rem;}

    [data-testid="stMetric"] {
        background: var(--secondary-background-color);
        border: 1px solid color-mix(
            in srgb, var(--text-color) 22%, transparent
        );
        border-radius: 12px; padding: 14px;
    }
    [data-testid="stMetricLabel"],
    [data-testid="stMetricValue"],
    [data-testid="stMetricDelta"] {
        color: var(--text-color) !important;
    }
    [data-testid="stMetric"] div,
    [data-testid="stMetric"] p {
        color: var(--text-color) !important;
    }

    .stCaption, [data-testid="stCaptionContainer"] p, .stMarkdown p {
        color: var(--text-color) !important;
        opacity: 0.85;
    }

    .photo-modal {
        display: none; position: fixed; inset: 0; z-index: 999999;
        background: rgba(3, 7, 18, .88); padding: 4vh 4vw;
        align-items: center; justify-content: center;
        backdrop-filter: blur(8px);
    }
    .photo-modal:target {display: flex;}
    .photo-modal img {
        max-width: 92vw; max-height: 88vh; object-fit: contain;
        border-radius: 14px; box-shadow: 0 24px 80px rgba(0,0,0,.65);
    }
    .photo-modal-close {
        position: absolute; top: 22px; right: 30px; color: #FFFFFF !important;
        font-size: 42px; line-height: 1; text-decoration: none !important;
    }
    .photo-details-grid {
        display: grid; grid-template-columns: 130px minmax(0, 1fr);
        column-gap: 24px; row-gap: 8px; align-items: start;
        line-height: 1.35; margin-top: 8px;
    }
    .photo-details-label {font-weight: 700; color: var(--text-color) !important;}
    .photo-details-value {overflow-wrap: anywhere; color: var(--text-color) !important;}

    .js-plotly-plot .plotly .legendtext,
    .js-plotly-plot .plotly .legendtitletext,
    .js-plotly-plot .plotly .gtitle,
    .js-plotly-plot .plotly .xtick text,
    .js-plotly-plot .plotly .ytick text,
    .js-plotly-plot .plotly .annotation-text,
    .js-plotly-plot .plotly .cbtitle text,
    .js-plotly-plot .plotly .colorbar text {
        fill: var(--text-color) !important;
        color: var(--text-color) !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# Shared visual system used by Incident, Evacuation, and future portal modules.
from evacuation.styles import apply_styles as apply_portal_styles
apply_portal_styles()


def handle_logout() -> None:
    """Purge all session-related credentials and restart the execution context."""
    try:
        supabase.auth.sign_out()
    except Exception:
        pass
    for key in (
        "admin_authenticated", "admin_user_id", "admin_email", "admin_username",
        "admin_role", "assigned_dashboard", "supabase_access_token",
        "supabase_refresh_token",
    ):
        st.session_state.pop(key, None)
    st.rerun()


def apply_chart_theme(figure):
    """Match Plotly to Streamlit's active theme and keep exports opaque."""
    try:
        theme_name = st.context.theme.type
    except (AttributeError, RuntimeError):
        theme_name = st.get_option("theme.base") or "light"

    is_dark = str(theme_name).lower() == "dark"
    background = "#0E1117" if is_dark else "#FFFFFF"
    foreground = "#F8FAFC" if is_dark else "#262730"
    grid = "rgba(248,250,252,0.16)" if is_dark else "rgba(38,39,48,0.14)"
    zero_line = "rgba(248,250,252,0.24)" if is_dark else "rgba(38,39,48,0.22)"

    figure.update_layout(
        paper_bgcolor=background,
        plot_bgcolor=background,
        font=dict(color=foreground, size=12),
        title_font=dict(color=foreground),
        legend=dict(
            font=dict(color=foreground, size=11),
            bgcolor=background,
        ),
        xaxis=dict(
            color=foreground,
            gridcolor=grid,
            zerolinecolor=zero_line,
        ),
        yaxis=dict(
            color=foreground,
            gridcolor=grid,
            zerolinecolor=zero_line,
        ),
    )
    return figure


def apply_pie_chart_theme(figure):
    """Apply the final adaptive donut design with readable slice labels."""
    figure = apply_chart_theme(figure)
    background = figure.layout.paper_bgcolor
    figure.update_traces(
        textfont=dict(size=12),
        insidetextfont=dict(color="#FFFFFF", size=12),
        marker=dict(line=dict(color=background, width=1.5)),
    )
    return figure


def connect_admin_db() -> sqlite3.Connection:
    """Open the local upload-history database."""
    connection = sqlite3.connect(ADMIN_DB_FILE, check_same_thread=False)
    connection.row_factory = sqlite3.Row
    return connection


def initialize_admin_db() -> None:
    """Create the local upload-history table; accounts live in Supabase."""
    with connect_admin_db() as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS upload_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                uploaded_at TEXT NOT NULL,
                username TEXT NOT NULL,
                performed_by_role TEXT,
                dashboard_name TEXT NOT NULL,
                filename TEXT NOT NULL,
                action TEXT NOT NULL,
                record_count INTEGER NOT NULL
            )
            """
        )
        history_columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(upload_history)")
        }
        if "performed_by_role" not in history_columns:
            connection.execute(
                "ALTER TABLE upload_history ADD COLUMN performed_by_role TEXT"
            )
        connection.execute(
            "UPDATE upload_history SET performed_by_role = 'unknown' "
            "WHERE performed_by_role IS NULL"
        )


def require_admin_client():
    """Return the protected Supabase client used only for account administration."""
    if supabase_admin is None:
        raise RuntimeError(
            "Account creation requires the Supabase secret/service-role key. "
            "Add secret_key under [supabase] in .streamlit/secrets.toml, "
            "then restart Streamlit."
        )
    return supabase_admin


def authenticate_admin(email: str, password: str) -> bool:
    """Authenticate with Supabase and load the active application account."""
    try:
        auth_response = supabase.auth.sign_in_with_password(
            {"email": email.strip(), "password": password}
        )
        user = auth_response.user
        session = auth_response.session
        if user is None or session is None:
            return False

        account_response = (
            supabase.table("accounts")
            .select("id,email,username,role,assigned_dashboard,active")
            .eq("id", str(user.id))
            .maybe_single()
            .execute()
        )
        account = account_response.data
        if not account or not account.get("active"):
            supabase.auth.sign_out()
            return False

        st.session_state.supabase_access_token = session.access_token
        st.session_state.supabase_refresh_token = session.refresh_token
        st.session_state.admin_user_id = str(user.id)
        st.session_state.admin_email = account["email"]
        st.session_state.admin_username = account["username"]
        st.session_state.admin_role = account["role"]
        st.session_state.assigned_dashboard = account.get("assigned_dashboard")
        return True
    except Exception:
        return False


def get_admin_role(username: str) -> str:
    """Return an active account's current role from Supabase."""
    response = (
        supabase.table("accounts")
        .select("role,active")
        .eq("username", username.strip())
        .maybe_single()
        .execute()
    )
    account = response.data
    return account["role"] if account and account.get("active") else "inactive"


def get_assigned_dashboard(username: str) -> str | None:
    """Return the dashboard assigned to an active Operator account."""
    response = (
        supabase.table("accounts")
        .select("assigned_dashboard")
        .eq("username", username.strip())
        .maybe_single()
        .execute()
    )
    return response.data.get("assigned_dashboard") if response.data else None


def change_admin_password(username: str, new_password: str) -> None:
    """Change the currently authenticated user's Supabase password."""
    supabase.auth.update_user({"password": new_password})


def create_admin_account(
    email: str,
    username: str,
    password: str,
    role: str,
    created_by: str,
    assigned_dashboard: str | None = None,
) -> None:
    """Create a Supabase Auth user and its application-role profile."""
    client = require_admin_client()
    creator_id = st.session_state.get("admin_user_id")
    auth_response = client.auth.admin.create_user(
        {
            "email": email.strip().lower(),
            "password": password,
            "email_confirm": True,
            "user_metadata": {"username": username.strip()},
        }
    )
    user = auth_response.user
    if user is None:
        raise RuntimeError("Supabase did not return the new user account.")

    client.table("accounts").upsert(
        {
            "id": str(user.id),
            "email": email.strip().lower(),
            "username": username.strip(),
            "role": role,
            "assigned_dashboard": assigned_dashboard if role == "operator" else None,
            "active": True,
            "created_by": creator_id,
            "updated_at": datetime.now().isoformat(),
        }
    ).execute()


def record_upload(
    username: str,
    dashboard_name: str,
    filename: str,
    action: str,
    record_count: int,
) -> None:
    """Append an auditable CSV/Excel event to upload history."""
    performed_by_role = get_admin_role(username)
    with connect_admin_db() as connection:
        connection.execute(
            """
            INSERT INTO upload_history
                (uploaded_at, username, performed_by_role, dashboard_name,
                 filename, action, record_count)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now().isoformat(timespec="seconds"),
                username,
                performed_by_role,
                dashboard_name,
                filename,
                action,
                record_count,
            ),
        )


def load_upload_history() -> pd.DataFrame:
    """Return the 200 latest upload events, including IDs used for selection."""
    with connect_admin_db() as connection:
        return pd.read_sql_query(
            """
            SELECT id AS "Log ID",
                   uploaded_at AS "Date and time",
                   username AS "Updated by",
                   CASE performed_by_role
                       WHEN 'super_admin' THEN 'Super Admin'
                       WHEN 'operator' THEN 'Operator'
                       WHEN 'admin' THEN 'Admin'
                       ELSE 'Unknown'
                   END AS "User role",
                   dashboard_name AS "Dashboard",
                   filename AS "Data file",
                   action AS "Action",
                   record_count AS "Records"
            FROM upload_history
            ORDER BY id DESC
            LIMIT 200
            """,
            connection,
        )


def load_admin_accounts() -> pd.DataFrame:
    """Load safe Supabase account details without authentication secrets."""
    response = (
        supabase.table("accounts")
        .select("username,email,role,assigned_dashboard,active,created_at,created_by")
        .order("username")
        .execute()
    )
    rows = response.data or []
    if not rows:
        return pd.DataFrame(columns=[
            "Admin username", "Email", "Role", "Dashboard access",
            "Status", "Created date", "Created by"
        ])
    dataframe = pd.DataFrame(rows).rename(columns={
        "username": "Admin username", "email": "Email",
        "assigned_dashboard": "Dashboard access", "active": "Status",
        "created_at": "Created date", "created_by": "Created by",
    })
    dataframe["Role"] = dataframe.pop("role").map({
        "super_admin": "Super Admin", "admin": "Admin", "operator": "Operator"
    }).fillna("Unknown")
    dataframe["Dashboard access"] = dataframe["Dashboard access"].fillna("All dashboards")
    dataframe["Status"] = dataframe["Status"].map({True: "Active", False: "Inactive"})
    return dataframe


def delete_upload_history(history_ids: list[int]) -> int:
    """Delete only the selected upload-history IDs and return the rows removed."""
    if not history_ids:
        return 0

    placeholders = ", ".join("?" for _ in history_ids)
    with connect_admin_db() as connection:
        cursor = connection.execute(
            f"DELETE FROM upload_history WHERE id IN ({placeholders})",
            tuple(history_ids),
        )
    return cursor.rowcount


def update_account_access(
    username: str, role: str, assigned_dashboard: str | None = None
) -> None:
    """Update a role and the one-dashboard assignment used by Operators."""
    if role not in {"operator", "admin", "super_admin"}:
        raise ValueError("Invalid admin role.")
    supabase.table("accounts").update({
        "role": role,
        "assigned_dashboard": assigned_dashboard if role == "operator" else None,
        "active": True,
        "updated_at": datetime.now().isoformat(),
    }).eq("username", username.strip()).execute()


@st.dialog("⚙️ Admin settings", width="large")
def show_admin_settings(
    admin_name: str, admin_role: str, available_dashboards: list[str]
) -> None:
    """Render the account, role, password, and history controls in a modal dialog."""
    operator_dashboards = [
        name for name in available_dashboards if name != OVERALL_DASHBOARD_NAME
    ]
    password_tab, account_tab, admin_list_tab, history_tab = st.tabs(
        ["Change password", "Add admin", "Admin list", "Upload history"]
    )

    with password_tab:
        with st.form("change_password_form", clear_on_submit=True):
            current_password = st.text_input("Current password", type="password")
            new_password = st.text_input("New password", type="password")
            confirm_password = st.text_input(
                "Confirm new password", type="password"
            )
            change_password_submitted = st.form_submit_button(
                "Change password", use_container_width=True
            )
        if change_password_submitted:
            if not authenticate_admin(
                st.session_state.get("admin_email", ""), current_password
            ):
                st.error("The current password is incorrect.")
            elif len(new_password) < 8:
                st.error("The new password must contain at least 8 characters.")
            elif new_password != confirm_password:
                st.error("The new passwords do not match.")
            else:
                change_admin_password(admin_name, new_password)
                st.success("Your password was changed successfully.")

    with account_tab:
        with st.form("add_admin_form", clear_on_submit=True):
            added_email = st.text_input("New account email")
            added_username = st.text_input("New admin username")
            if admin_role == "super_admin":
                selected_role_label = st.selectbox(
                    "Account role", ["Operator", "Admin", "Super Admin"]
                )
            else:
                selected_role_label = st.selectbox(
                    "Account role", ["Operator", "Admin"]
                )
                st.caption("Admins can create Operator or Admin accounts.")
            assigned_dashboard = st.selectbox(
                "Operator dashboard",
                operator_dashboards if operator_dashboards else ["No Dashboards Available"],
                help="Used only when the new account role is Operator.",
            )
            added_password = st.text_input("New admin password", type="password")
            added_password_confirm = st.text_input(
                "Confirm admin password", type="password"
            )
            add_admin_submitted = st.form_submit_button(
                "Add admin account", use_container_width=True
            )
        if add_admin_submitted:
            if "@" not in added_email or not added_email.strip():
                st.error("Enter a valid email address for the new account.")
            elif not added_username.strip():
                st.error("Enter a username for the new admin.")
            elif len(added_password) < 8:
                st.error("The password must contain at least 8 characters.")
            elif added_password != added_password_confirm:
                st.error("The passwords do not match.")
            else:
                try:
                    selected_role = (
                        "super_admin"
                        if selected_role_label == "Super Admin"
                        else "operator" if selected_role_label == "Operator" else "admin"
                    )
                    create_admin_account(
                        added_email,
                        added_username,
                        added_password,
                        selected_role,
                        admin_name,
                        assigned_dashboard if assigned_dashboard != "No Dashboards Available" else None,
                    )
                except Exception as error:
                    st.error(f"The account could not be created: {error}")
                else:
                    st.success(
                        f"Admin account '{added_username.strip()}' was created."
                    )

    with admin_list_tab:
        admin_accounts = load_admin_accounts()
        st.caption(f"Total admin accounts: {len(admin_accounts):,}")
        st.dataframe(
            admin_accounts,
            hide_index=True,
            use_container_width=True,
        )
        if admin_role in {"admin", "super_admin"}:
            st.divider()
            st.markdown("#### Change account role")
            editable_accounts = [
                account_row["Admin username"]
                for _, account_row in admin_accounts.iterrows()
                if account_row["Admin username"].casefold() != admin_name.casefold()
                and (
                    admin_role == "super_admin"
                    or account_row["Role"] in {"Operator", "Admin"}
                )
            ]
            if editable_accounts:
                with st.form("change_admin_role_form"):
                    role_username = st.selectbox(
                        "Admin account", editable_accounts
                    )
                    current_account = admin_accounts.loc[
                        admin_accounts["Admin username"].str.casefold()
                        == role_username.casefold(),
                    ].iloc[0]
                    current_role_label = current_account["Role"]
                    role_choices = (
                        ["Operator", "Admin", "Super Admin"]
                        if admin_role == "super_admin"
                        else ["Operator", "Admin"]
                    )
                    new_role_label = st.selectbox(
                        "New role",
                        role_choices,
                        index=(role_choices.index(current_role_label)
                               if current_role_label in role_choices else 0),
                    )
                    operator_dashboard = st.selectbox(
                        "Operator dashboard access",
                        operator_dashboards if operator_dashboards else ["No Dashboards Available"],
                        index=(operator_dashboards.index(current_account["Dashboard access"])
                               if current_account["Dashboard access"] in operator_dashboards else 0),
                        help="Used only when the selected role is Operator.",
                    )
                    confirm_role_change = st.checkbox(
                        "I confirm this account role change"
                    )
                    role_change_submitted = st.form_submit_button(
                        "Update account role", use_container_width=True
                    )
                if role_change_submitted:
                    if not confirm_role_change:
                        st.error("Confirm the role change before updating the account.")
                    else:
                        new_role = (
                            "super_admin"
                            if new_role_label == "Super Admin"
                            else "operator" if new_role_label == "Operator" else "admin"
                        )
                        update_account_access(
                            role_username,
                            new_role,
                            operator_dashboard if operator_dashboard != "No Dashboards Available" else None,
                        )
                        st.success(
                            f"{role_username} is now assigned as {new_role_label}."
                        )
                        st.rerun()
            else:
                st.caption("Add another admin account before changing roles.")
            st.caption("For security, you cannot change your own role.")
        else:
            st.caption("Only an Admin or Super Admin can manage account access.")

    with history_tab:
        history = load_upload_history()
        if history.empty:
            st.caption("No CSV or Excel uploads have been recorded yet.")
        elif admin_role == "super_admin":
            selectable_history = history.copy()
            selectable_history.insert(0, "Select", False)
            edited_history = st.data_editor(
                selectable_history,
                hide_index=True,
                use_container_width=True,
                disabled=list(history.columns),
                column_config={
                    "Select": st.column_config.CheckboxColumn(
                        "Select",
                        help="Check each history record you want to delete.",
                        default=False,
                    ),
                    "Log ID": None,
                },
                key="upload_history_selector",
            )
            selected_history_ids = edited_history.loc[
                edited_history["Select"], "Log ID"
            ].astype(int).tolist()

            st.divider()
            st.warning("Selected upload-history records will be permanently deleted.")
            confirm_delete_history = st.checkbox(
                "I confirm that I want to delete the selected history record(s)",
                key="confirm_selected_history_deletion",
            )
            delete_history_submitted = st.button(
                f"🗑️ Delete selected history ({len(selected_history_ids)})",
                use_container_width=True,
                disabled=not selected_history_ids,
            )
            if delete_history_submitted:
                if not confirm_delete_history:
                    st.error("Confirm the deletion before deleting the selected records.")
                else:
                    deleted_count = delete_upload_history(selected_history_ids)
                    st.success(
                        f"{deleted_count:,} selected upload-history record(s) deleted."
                    )
                    st.rerun()
        else:
            st.dataframe(
                history,
                hide_index=True,
                use_container_width=True,
                column_config={"Log ID": None},
            )
            st.caption("Only a Super Admin can delete upload-history records.")


initialize_admin_db()


def parse_dashboard_dates(values: pd.Series) -> pd.Series:
    """Parse dashboard dates as day/month/year and reject implausible years."""
    raw_text = values.astype("string").str.strip()
    normalized = raw_text.copy()

    short_day_month = raw_text.str.fullmatch(r"\d{1,2}[-/.]\d{1,2}", na=False)
    if short_day_month.any():
        normalized.loc[short_day_month] = (
            raw_text.loc[short_day_month].str.replace(r"[-/.]", "/", regex=True)
            + f"/{datetime.now().year}"
        )

    try:
        parsed = pd.to_datetime(
            normalized,
            format="mixed",
            dayfirst=True,
            errors="coerce",
        )
    except (TypeError, ValueError):
        parsed = pd.to_datetime(normalized, dayfirst=True, errors="coerce")

    parsed = pd.Series(parsed, index=values.index)

    end_of_today = pd.Timestamp.now().normalize() + pd.Timedelta(days=1)
    swapped_month_day = pd.to_datetime(
        {
            "year": parsed.dt.year,
            "month": parsed.dt.day,
            "day": parsed.dt.month,
        },
        errors="coerce",
    )
    swap_mask = (
        parsed.notna()
        & (parsed >= end_of_today)
        & swapped_month_day.notna()
        & (swapped_month_day < end_of_today)
    )
    parsed = parsed.where(~swap_mask, swapped_month_day)

    implausible = parsed.notna() & (
        (parsed.dt.year < 1900) | (parsed >= end_of_today)
    )
    return parsed.where(~implausible, pd.NaT)


def clean_columns(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Normalize known column aliases, dates, and coordinates for dashboard use."""
    dataframe = dataframe.copy()
    dataframe.columns = dataframe.columns.astype(str).str.strip()
    aliases = {
        "Spot Name": "Spot name",
        "Latitude": "latitude",
        "Longitude": "longitude",
        "Register Date and time": "Registered Date",
        "Register Date": "Registered Date",
        "Updated Date and Time": "Updated date and time",
        "Type of incident": "Type of Incident",
        "Incident Type": "Type of Incident",
        "Traffic Condition": "Traffic",
        "Traffic condition": "Traffic",
        "Traffic Status": "Traffic",
        "City/ Municipality": "Municipality",
        "City/Municipality": "Municipality",
        "City / Municipality": "Municipality",
        "City or Municipality": "Municipality",
        "City Municipality": "Municipality",
        "Municipality/City": "Municipality",
        "Municipality Name": "Municipality",
        "Municipality name": "Municipality",
        "municipality": "Municipality",
        "MUNICIPALITY": "Municipality",
        "incident_date": "Registered Date",
        "dashboard_name": "Source Dashboard",
        "service": "Type of Incident",
        "traffic": "Traffic",
        "registrant": "Registrant",
        "photo_url": "Photo",
        "barangay": "Barangay",
    }
    dataframe = dataframe.rename(columns={k: v for k, v in aliases.items() if k in dataframe})

    for column in ("Registered Date", "Updated date and time"):
        if column in dataframe:
            dataframe[column] = parse_dashboard_dates(dataframe[column])

    for column in ("latitude", "longitude"):
        if column in dataframe:
            dataframe[column] = pd.to_numeric(dataframe[column], errors="coerce")
    return dataframe


@st.cache_data(show_spinner=False)
def load_data_file(source, filename: str | None = None) -> pd.DataFrame:
    """Read and clean CSV, XLSX, or XLS dashboard records."""
    source_name = filename or getattr(source, "name", str(source))
    extension = Path(source_name).suffix.lower()

    if extension in {".xlsx", ".xls"}:
        if hasattr(source, "seek"):
            source.seek(0)
        try:
            with pd.ExcelFile(source) as workbook:
                normalized_sheet_names = {
                    "".join(
                        character
                        for character in sheet_name.strip().casefold()
                        if character.isalnum()
                    ): sheet_name
                    for sheet_name in workbook.sheet_names
                }
                data_sheet = normalized_sheet_names.get("spotlist")
                if data_sheet is None:
                    data_sheet = next(
                        (
                            sheet_name
                            for sheet_name in workbook.sheet_names
                            if "photo" not in sheet_name.casefold()
                        ),
                        workbook.sheet_names[0],
                    )

                preview = pd.read_excel(
                    workbook, sheet_name=data_sheet, header=None, nrows=20
                )
                known_headings = {
                    "id",
                    "no",
                    "number",
                    "spot name",
                    "spotname",
                    "address",
                    "latitude",
                    "longitude",
                    "registrant",
                    "register date and time",
                    "registered date",
                    "description",
                    "incident details",
                    "barangay",
                    "municipality",
                }
                best_header_row = 0
                best_header_score = 0
                for row_number, row in preview.iterrows():
                    row_values = {
                        " ".join(str(value).strip().casefold().split())
                        for value in row.dropna().tolist()
                    }
                    header_score = len(row_values & known_headings)
                    if header_score > best_header_score:
                        best_header_score = header_score
                        best_header_row = int(row_number)

                dataframe = pd.read_excel(
                    workbook,
                    sheet_name=data_sheet,
                    header=best_header_row if best_header_score >= 2 else 0,
                )
                return clean_columns(dataframe)
        except ImportError as error:
            package = "openpyxl" if extension == ".xlsx" else "xlrd"
            raise RuntimeError(
                f"Excel support requires '{package}'. Install it with: "
                f"python -m pip install {package}"
            ) from error

    decoding_error = None
    for encoding in ("utf-8-sig", "cp932", "shift_jis", "cp1252", "iso-8859-1"):
        if hasattr(source, "seek"):
            source.seek(0)
        try:
            return clean_columns(pd.read_csv(source, encoding=encoding))
        except UnicodeDecodeError as error:
            decoding_error = error
    raise RuntimeError("The CSV text encoding could not be detected.") from decoding_error


def data_file_mime(filename: str) -> str:
    """Return an appropriate download MIME type for CSV or Excel data."""
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if extension == ".xls":
        return "application/vnd.ms-excel"
    return "text/csv"


def gallery_caption(dataframe: pd.DataFrame, row_index: int) -> str:
    """Build an image caption with incident identity and photo description."""
    if row_index < 0 or row_index >= len(dataframe):
        return "Incident photo"
    row = dataframe.iloc[row_index]
    parts = []
    for column in ("ID", "Spot name", "Barangay", "Municipality"):
        if column in dataframe and pd.notna(row.get(column)):
            value = str(row.get(column)).strip()
            if value and value.casefold() not in {"nan", "none"}:
                parts.append(value)
    caption = " · ".join(parts[:3]) or f"Incident record {row_index + 1}"
    description_headings = {
        "description",
        "photo description",
        "incident description",
        "incident details",
        "details",
        "remarks",
        "title",
    }
    for column in dataframe.columns:
        if str(column).strip().casefold() not in description_headings:
            continue
        value = row.get(column)
        if pd.notna(value):
            description = str(value).strip()
            if description and description.casefold() not in {"nan", "none"}:
                return f"{caption}\nDescription: {description[:180]}"
    return caption


def gallery_image_source(image) -> str:
    """Return a browser-readable URL for a gallery URL or embedded image."""
    if isinstance(image, str):
        return image
    image_bytes = bytes(image)
    if image_bytes.startswith(b"\xff\xd8\xff"):
        mime_type = "image/jpeg"
    elif image_bytes.startswith((b"GIF87a", b"GIF89a")):
        mime_type = "image/gif"
    elif image_bytes.startswith(b"RIFF") and image_bytes[8:12] == b"WEBP":
        mime_type = "image/webp"
    else:
        mime_type = "image/png"
    encoded_image = base64.b64encode(image_bytes).decode("ascii")
    return f"data:{mime_type};base64,{encoded_image}"


def worksheet_photo_metadata(worksheet, anchor_row: int, anchor_column: int) -> dict:
    """Read the metadata block immediately above an embedded worksheet photo."""
    image_row = int(anchor_row) + 1
    metadata = {}
    accepted_labels = {
        "no": "No",
        "number": "No",
        "spotname": "SpotName",
        "spot name": "SpotName",
        "name": "Name",
        "title": "Title",
        "description": "Description",
        "photo description": "Description",
        "incident description": "Description",
        "incident details": "Description",
        "details": "Description",
        "remarks": "Description",
        "address": "Address",
        "date": "Date",
    }

    anchor_excel_column = int(anchor_column) + 1
    candidate_columns = range(
        max(anchor_excel_column - 1, 1),
        min(anchor_excel_column + 2, worksheet.max_column) + 1,
    )
    label_column = None
    best_label_count = 0
    for candidate_column in candidate_columns:
        label_count = 0
        for row_number in range(image_row - 1, max(image_row - 12, 0), -1):
            candidate_label = worksheet.cell(
                row=row_number, column=candidate_column
            ).value
            normalized_candidate = " ".join(
                str(candidate_label or "").strip().casefold().split()
            )
            if normalized_candidate in accepted_labels:
                label_count += 1
        if label_count > best_label_count:
            best_label_count = label_count
            label_column = candidate_column
    if label_column is None:
        return {}
    value_column = label_column + 1

    for row_number in range(image_row - 1, max(image_row - 12, 0), -1):
        label = worksheet.cell(row=row_number, column=label_column).value
        value = worksheet.cell(row=row_number, column=value_column).value
        normalized_label = " ".join(str(label or "").strip().casefold().split())
        metadata_key = accepted_labels.get(normalized_label)
        if not metadata_key or value is None:
            continue
        if isinstance(value, (pd.Timestamp, datetime, date)):
            value_text = pd.Timestamp(value).strftime("%Y-%m-%d %H:%M:%S")
        else:
            value_text = str(value).strip()
        if value_text and value_text.casefold() not in {"nan", "none"}:
            metadata[metadata_key] = value_text
    return metadata


def worksheet_photo_caption(worksheet, anchor_row: int, anchor_column: int) -> str:
    """Build a short caption from an embedded photo's worksheet metadata."""
    metadata = worksheet_photo_metadata(worksheet, anchor_row, anchor_column)

    parts = []
    for key in ("SpotName", "No", "Date"):
        if metadata.get(key):
            parts.append(metadata[key])
    caption = " · ".join(parts) or f"{worksheet.title} photo"
    description = metadata.get("Description") or metadata.get("Title")
    if not description and metadata.get("Name", "").casefold() != "photos":
        description = metadata.get("Name")
    if description:
        caption = f"{caption}\nDescription: {description[:180]}"
    return caption


def dataframe_photo_details(dataframe: pd.DataFrame, row_index: int) -> dict:
    """Return Excel-style gallery details from a normal incident-data row."""
    if row_index < 0 or row_index >= len(dataframe):
        return {}
    row = dataframe.iloc[row_index]
    aliases = {
        "No": ("no", "number", "id"),
        "SpotName": ("spotname", "spot name"),
        "Name": ("name", "photo name"),
        "Title": ("title",),
        "Address": ("address", "location"),
        "Date": ("date", "registered date", "register date and time", "registered date and time"),
    }
    normalized_columns = {
        " ".join(str(column).strip().casefold().split()): column
        for column in dataframe.columns
    }
    details = {}
    for label, candidates in aliases.items():
        for candidate in candidates:
            column = normalized_columns.get(candidate)
            if column is None or pd.isna(row.get(column)):
                continue
            value = row.get(column)
            if isinstance(value, (pd.Timestamp, datetime, date)):
                value_text = pd.Timestamp(value).strftime("%d %b %Y %H:%M:%S")
            else:
                value_text = str(value).strip()
            if value_text and value_text.casefold() not in {"nan", "none"}:
                details[label] = value_text
                break
    return details


def spot_list_photo_address(dataframe: pd.DataFrame, photo_details: dict) -> str:
    """Find a photo's address in the uploaded Spot list data."""
    if dataframe.empty:
        return ""
    normalized_columns = {
        " ".join(str(column).strip().casefold().split()): column
        for column in dataframe.columns
    }
    match_groups = (
        (photo_details.get("No"), ("no", "number", "id")),
        (photo_details.get("SpotName"), ("spotname", "spot name")),
    )
    matched_row = None
    for expected_value, candidate_columns in match_groups:
        if not expected_value:
            continue
        expected_text = " ".join(str(expected_value).strip().casefold().split())
        for candidate in candidate_columns:
            column = normalized_columns.get(candidate)
            if column is None:
                continue
            normalized_values = dataframe[column].fillna("").astype(str).map(
                lambda value: " ".join(value.strip().casefold().split())
            )
            matches = dataframe.loc[normalized_values == expected_text]
            if not matches.empty:
                matched_row = matches.iloc[0]
                break
        if matched_row is not None:
            break
    if matched_row is None:
        return ""

    address_columns = ("address", "location")
    for heading in address_columns:
        column = normalized_columns.get(heading)
        if column is None or pd.isna(matched_row.get(column)):
            continue
        address = str(matched_row.get(column)).strip()
        if address and address.casefold() not in {"nan", "none"}:
            return address
    return ""


@st.cache_data(show_spinner=False)
def extract_gallery_items(
    file_bytes: bytes, filename: str, dataframe: pd.DataFrame
) -> list[dict]:
    """Extract URL/data-URI photos and embedded XLSX worksheet images."""
    gallery_items = []
    photo_headings = {
        "photo",
        "image",
        "picture",
        "photo url",
        "photo_url",
        "image url",
        "picture url",
        "attachment",
        "attachment url",
    }
    photo_columns = [
        column
        for column in dataframe.columns
        if str(column).strip().casefold() in photo_headings
    ]
    for row_index, row in dataframe.iterrows():
        for column in photo_columns:
            if pd.isna(row.get(column)):
                continue
            image_reference = str(row.get(column)).strip()
            if image_reference.startswith(("https://", "http://", "data:image/")):
                gallery_items.append(
                    {
                        "image": image_reference,
                        "caption": gallery_caption(dataframe, row_index),
                        "details": dataframe_photo_details(dataframe, row_index),
                    }
                )

    if Path(filename).suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(file_bytes), data_only=True)
            first_worksheet = workbook.worksheets[0]
            for worksheet in workbook.worksheets:
                for worksheet_image in getattr(worksheet, "_images", []):
                    try:
                        anchor = getattr(worksheet_image, "anchor", None)
                        anchor_from = getattr(anchor, "_from", None)
                        worksheet_row = int(getattr(anchor_from, "row", 1))
                        worksheet_column = int(getattr(anchor_from, "col", 0))
                        caption = worksheet_photo_caption(
                            worksheet, worksheet_row, worksheet_column
                        )
                        details = worksheet_photo_metadata(
                            worksheet, worksheet_row, worksheet_column
                        )
                        spot_list_address = spot_list_photo_address(
                            dataframe, details
                        )
                        if spot_list_address:
                            details["Address"] = spot_list_address
                        if (
                            caption == f"{worksheet.title} photo"
                            and worksheet is first_worksheet
                        ):
                            dataframe_row = max(worksheet_row - 1, 0)
                            caption = gallery_caption(dataframe, dataframe_row)
                            details = dataframe_photo_details(dataframe, dataframe_row)
                        gallery_items.append(
                            {
                                "image": worksheet_image._data(),
                                "caption": caption,
                                "details": details,
                            }
                        )
                    except Exception:
                        continue
            workbook.close()
        except ImportError:
            pass
        except Exception:
            pass

    return gallery_items


def normalize_category(value) -> str:
    """Normalize spacing/capitalization so case variants count as one category."""
    if pd.isna(value):
        return "Not Specified"
    text_val = " ".join(str(value).strip().split())
    return text_val.title() if text_val else "Not Specified"


def options_for(dataframe: pd.DataFrame, column: str) -> list[str]:
    """Build sorted, case-normalized, duplicate-free options for a filter widget."""
    if column not in dataframe:
        return []
    return sorted(
        dataframe[column]
        .dropna()
        .map(normalize_category)
        .loc[lambda values: values.ne("")]
        .unique()
        .tolist()
    )


def first_available_column(dataframe: pd.DataFrame, choices: list[str]) -> str | None:
    """Return the first candidate column that exists and contains usable values."""
    return next(
        (
            column
            for column in choices
            if column in dataframe and dataframe[column].notna().any()
        ),
        None,
    )


def add_source_label(dataframe: pd.DataFrame, dashboard_name: str) -> pd.DataFrame:
    """Add a source-dashboard label before combining files in Overall View."""
    dataframe = dataframe.copy()
    dataframe["Source Dashboard"] = dashboard_name
    return dataframe[["Source Dashboard", *[c for c in dataframe if c != "Source Dashboard"]]]


def category_counts(dataframe: pd.DataFrame, column: str, limit: int = 10) -> pd.DataFrame:
    """Count normalized categories while keeping Overall View file sources separate."""
    count_data = pd.DataFrame(
        {"Category": dataframe[column].map(normalize_category)}
    )
    if "Source Dashboard" in dataframe:
        count_data["Source Dashboard"] = dataframe["Source Dashboard"].map(
            normalize_category
        )
        count_data["Category"] = (
            count_data["Source Dashboard"] + " — " + count_data["Category"]
        )

    return (
        count_data.value_counts("Category")
        .head(limit)
        .rename("Incidents")
        .reset_index()
    )


def shorten_category(value: str, max_length: int = 20) -> str:
    """Shorten chart-only labels without altering the uploaded source data."""
    text_val = str(value).strip()
    return text_val if len(text_val) <= max_length else f"{text_val[: max_length - 1].rstrip()}…"


# ------------------------------------------------------------------------------
# Session State Initialization & Authentication
# ------------------------------------------------------------------------------
if "uploaded_dashboards" not in st.session_state:
    st.session_state.uploaded_dashboards = {}
if "admin_authenticated" not in st.session_state:
    st.session_state.admin_authenticated = False
if "admin_username" not in st.session_state:
    st.session_state.admin_username = None
if "admin_role" not in st.session_state:
    st.session_state.admin_role = None
if "admin_email" not in st.session_state:
    st.session_state.admin_email = None
if "admin_user_id" not in st.session_state:
    st.session_state.admin_user_id = None
if "assigned_dashboard" not in st.session_state:
    st.session_state.assigned_dashboard = None

# Restore the signed-in Supabase session after every Streamlit rerun.
if (
    st.session_state.get("supabase_access_token")
    and st.session_state.get("supabase_refresh_token")
):
    try:
        restored_session = supabase.auth.set_session(
            st.session_state.supabase_access_token,
            st.session_state.supabase_refresh_token,
        )
        if restored_session.session:
            st.session_state.supabase_access_token = restored_session.session.access_token
            st.session_state.supabase_refresh_token = restored_session.session.refresh_token
    except Exception:
        st.session_state.admin_authenticated = False

if not st.session_state.admin_authenticated:
    st.markdown(
        """
        <style>
        [data-testid="stAppViewContainer"] {
            background:
                radial-gradient(circle at 15% 20%, rgba(255, 75, 85, .18), transparent 30%),
                radial-gradient(circle at 85% 80%, rgba(52, 152, 219, .14), transparent 32%),
                linear-gradient(135deg, #090d14 0%, #111827 52%, #0b1220 100%);
        }
        [data-testid="stHeader"], [data-testid="stSidebar"] {display: none;}
        .block-container {
            max-width: 1180px;
            padding-top: 6vh;
            padding-bottom: 4rem;
        }
        .login-brand {
            text-align: center;
            margin-bottom: 1.4rem;
        }
        .login-brand h1 {
            margin: 0;
            color: #f8fafc;
            font-size: clamp(2rem, 4vw, 3rem);
            letter-spacing: -.04em;
        }
        .login-brand p {
            margin: .55rem 0 0;
            color: #94a3b8;
            font-size: 1rem;
        }
        div[data-testid="stForm"] {
            padding: 1.8rem 2rem 2rem;
            border: 1px solid rgba(148, 163, 184, .20);
            border-radius: 22px;
            background: rgba(17, 24, 39, .82);
            box-shadow: 0 24px 70px rgba(0, 0, 0, .35);
            backdrop-filter: blur(16px);
        }
        div[data-testid="stForm"] label,
        div[data-testid="stForm"] p,
        div[data-testid="stForm"] h1,
        div[data-testid="stForm"] h2,
        div[data-testid="stForm"] h3 {
            color: #f8fafc !important;
        }
        div[data-testid="stTextInput"] input {
            min-height: 48px;
            border-radius: 12px;
            border-color: #334155;
            background: #0f172a;
            color: #f8fafc !important;
            caret-color: #f8fafc;
        }
        div[data-testid="stTextInput"] input::placeholder {
            color: #94a3b8 !important;
        }
        div[data-testid="stFormSubmitButton"] button {
            min-height: 50px;
            margin-top: .65rem;
            border: 0;
            border-radius: 12px;
            background: linear-gradient(90deg, #ff4b55, #ff6570);
            box-shadow: 0 10px 24px rgba(255, 75, 85, .24);
            font-weight: 700;
            color: #ffffff !important;
        }
        .login-security {
            margin-top: 1.1rem;
            text-align: center;
            color: #94a3b8;
            font-size: .82rem;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    _, login_column, _ = st.columns([1, 1.15, 1])
    with login_column:
        st.markdown(
            """
            <div class="login-brand">
                <h1>SOCOCA Dashboard</h1>
                <p>Secure incident monitoring and operations portal</p>
            </div>
            """,
            unsafe_allow_html=True,
        )
        with st.form("account_login_form"):
            st.markdown("### Welcome back")
            st.caption("Sign in with your authorized account to continue.")
            login_email = st.text_input(
                "Email address", placeholder="Enter your account email"
            )
            login_password = st.text_input(
                "Password", type="password", placeholder="Enter your password"
            )
            login_submitted = st.form_submit_button(
                "Sign in securely  →", type="primary", use_container_width=True
            )
        st.markdown(
            '<div class="login-security">🔒 Protected role-based dashboard access</div>',
            unsafe_allow_html=True,
        )
    if login_submitted:
        if authenticate_admin(login_email, login_password):
            st.session_state.admin_authenticated = True
            st.rerun()
        else:
            with login_column:
                st.error("Incorrect email/password, or the account is inactive.")
    st.stop()

# One authenticated portal controls access to every monitoring module.
with st.sidebar:
    st.markdown(
        """
        <div class="brand">
            <span class="brand-icon">✦</span>
            <div>SOCOCA MONITORING<small>Unified Operations Portal</small></div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    portal_role_label = {
        "super_admin": "Super Admin",
        "admin": "Admin",
        "operator": "Operator",
    }.get(st.session_state.admin_role, "User")
    st.markdown(
        f"""
        <div class="user-chip">
            <span>👤</span>
            <span><small>{html.escape(portal_role_label)}</small><br>
            {html.escape(str(st.session_state.admin_username))}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if st.button(
        "Log out",
        icon=":material/logout:",
        use_container_width=True,
        key="portal_logout",
    ):
        handle_logout()
    st.divider()
    portal_section = st.radio(
        "Main menu",
        ["Incident Dashboard", "Evacuation Dashboard"],
        key="portal_section",
    )

if portal_section == "Evacuation Dashboard":
    # The evacuation module inherits the authenticated Incident account.
    import importlib
    import sys

    st.session_state.authenticated = True
    module_name = "evacuation.app"
    if module_name in sys.modules:
        importlib.reload(sys.modules[module_name])
    else:
        importlib.import_module(module_name)
    st.stop()

if "pending_dashboard" in st.session_state:
    st.session_state.selected_dashboard = st.session_state.pop("pending_dashboard")

# Fetch current databases dynamic references
db_dashboards = []
try:
    db_df = load_dashboard_data()
    if not db_df.empty and "dashboard_name" in db_df.columns:
        db_dashboards = db_df["dashboard_name"].dropna().unique().tolist()
except Exception:
    pass

all_dashboard_names = list(
    dict.fromkeys([
        OVERALL_DASHBOARD_NAME,
        *db_dashboards,
        *st.session_state.uploaded_dashboards.keys(),
    ])
)

if "selected_dashboard" not in st.session_state or st.session_state.selected_dashboard not in all_dashboard_names:
    st.session_state.selected_dashboard = all_dashboard_names[0]

account_name = st.session_state.admin_username
account_role = st.session_state.admin_role
assigned_dashboard = st.session_state.assigned_dashboard
if account_role == "operator":
    dashboard_names = [assigned_dashboard] if assigned_dashboard in all_dashboard_names else []
    if not dashboard_names:
        st.error(
            "Your assigned dashboard is not currently available. "
            "Ask an Admin or Super Admin to update your dashboard access."
        )
        if st.button("Log out"):
            handle_logout()
        st.stop()
    if st.session_state.selected_dashboard not in dashboard_names:
        st.session_state.selected_dashboard = assigned_dashboard
else:
    dashboard_names = all_dashboard_names

# ------------------------------------------------------------------------------
# Sidebar Navigation
# ------------------------------------------------------------------------------
with st.sidebar:
    st.title(" Dashboards")
    selected_dashboard = st.radio(
        "Choose a dashboard",
        dashboard_names,
        label_visibility="collapsed",
        key="selected_dashboard",
    )
    if account_role == "operator":
        st.caption(f"Operator access: {assigned_dashboard} only.")
    else:
        st.caption("Select a location to display its uploaded-data dashboard.")

    # Fetch mapping of stored files from Supabase Storage
    storage_files_map = fetch_uploaded_files_from_supabase()

    # Generate file dynamic list based on verified state and database tables
    available_dashboard_names = [
        name for name in all_dashboard_names if name != OVERALL_DASHBOARD_NAME
    ]

    if account_role == "operator":
        available_dashboard_names = [
            name for name in available_dashboard_names if name == assigned_dashboard
        ]
    elif selected_dashboard != OVERALL_DASHBOARD_NAME:
        available_dashboard_names = [
            name for name in available_dashboard_names if name == selected_dashboard
        ]

    total_data_files = len(available_dashboard_names)
    st.divider()
    st.markdown(f"#### Available data files ({total_data_files})")

    for number, name in enumerate(available_dashboard_names, start=1):
        item = st.session_state.uploaded_dashboards.get(name)
        if item:
            file_name = item["filename"]
            file_data = item["data"]
            mime_type = data_file_mime(file_name)
        else:
            file_name = storage_files_map.get(name, f"{name}_data.csv")
            dash_df = load_dashboard_data(name)
            file_data = dash_df.to_csv(index=False).encode("utf-8-sig")
            mime_type = data_file_mime(file_name)

        st.markdown(f"**{number}. {name}**  \n`{file_name}`")

        can_delete_dashboard = account_role in {"admin", "super_admin"}
        btn_col1, btn_col2 = (
            st.columns([3, 1]) if can_delete_dashboard else (st.container(), None)
        )
        with btn_col1:
            st.download_button(
                "Download",
                data=file_data,
                file_name=file_name,
                mime=mime_type,
                icon=":material/download:",
                key=f"download_uploaded_{number}_{name}",
                use_container_width=True,
            )
        if can_delete_dashboard and btn_col2 is not None:
            with btn_col2:
                if st.button(
                    " ",
                    icon=":material/delete:",
                    key=f"delete_uploaded_{number}_{name}",
                    help=f"Delete {name} dataset",
                    use_container_width=True,
                ):
                    if delete_dashboard_data_from_supabase(name, account_name):
                        st.toast(f"Data file '{name}' was deleted.")
                        st.rerun()

left, right = st.columns([3, 1], gap="large")
with right:
    admin_name = account_name
    admin_role = account_role
    role_label = {
        "super_admin": "Super Admin",
        "admin": "Admin",
        "operator": "Operator",
    }.get(admin_role, "Operator")
    st.subheader("Update data" if admin_role == "operator" else "Upload data")
    if admin_role in {"operator", "admin", "super_admin"}:
        with st.form("add_dashboard_form", clear_on_submit=True):
            uploaded_file = st.file_uploader(
                "Choose a CSV or Excel file", type=["csv", "xlsx", "xls"]
            )
            if admin_role == "operator":
                new_dashboard_name = st.text_input(
                    "Assigned dashboard",
                    value=assigned_dashboard,
                    disabled=True,
                )
                st.caption("Operators can update only their assigned dashboard.")
            else:
                new_dashboard_name = st.text_input(
                    "Dashboard name",
                    placeholder="Example: Laguna",
                )
            confirmed = st.checkbox(
                "I confirm that I want to add or update this data file"
            )
            add_dashboard = st.form_submit_button(
                "⬆️ Update assigned dashboard"
                if admin_role == "operator"
                else "➕ Add or update dashboard",
                type="primary",
                use_container_width=True,
            )

        if add_dashboard:
            clean_name = new_dashboard_name.strip()
            if uploaded_file is None:
                st.error("Please choose a CSV or Excel file.")
            elif not clean_name:
                st.error("Please enter a dashboard name.")
            elif (
                admin_role == "operator"
                and clean_name.casefold() != assigned_dashboard.casefold()
            ):
                st.error("Operators can update only their assigned dashboard.")
            elif clean_name.casefold() == OVERALL_DASHBOARD_NAME.casefold():
                st.error("Overall View is generated automatically and cannot be uploaded.")
            elif not confirmed:
                st.warning("Please confirm before adding the data file.")
            else:
                uploaded_bytes = uploaded_file.getvalue()

                # --- SUPABASE STORAGE UPLOAD & FEEDBACK ---
                target_filename = f"{clean_name}_{uploaded_file.name}"
                public_url = upload_photo_to_supabase(uploaded_bytes, target_filename)

                if not public_url:
                    st.error("Failed to upload file to Supabase. Check console logs or RLS policies.")
                else:
                    try:
                        uploaded_df = load_data_file(
                            BytesIO(uploaded_bytes), uploaded_file.name
                        )
                        if uploaded_df.empty:
                            raise ValueError("The uploaded file contains no records.")
                        uploaded_gallery = extract_gallery_items(
                            uploaded_bytes, uploaded_file.name, uploaded_df
                        )
                    except Exception as error:
                        st.error(f"The data file could not be added: {error}")
                    else:
                        existing_name = next(
                            (
                                name
                                for name in st.session_state.uploaded_dashboards
                                if name.casefold() == clean_name.casefold()
                            ),
                            None,
                        )
                        new_item = {
                            "filename": uploaded_file.name,
                            "data": uploaded_bytes,
                            "gallery": uploaded_gallery,
                            "public_url": public_url,
                        }

                        if existing_name:
                            st.session_state.uploaded_dashboards[existing_name] = new_item
                            target_name = existing_name
                            upload_action = "Updated"
                        elif admin_role == "operator":
                            st.error("Your assigned dashboard is not available to update.")
                            st.stop()
                        else:
                            st.session_state.uploaded_dashboards[clean_name] = new_item
                            target_name = clean_name
                            upload_action = "Created"

                        # Database Persistence Step
                        try:
                            save_dataframe_to_supabase(uploaded_df, target_name)
                        except Exception as db_err:
                            st.warning(f"Saved locally, but Supabase insert failed: {db_err}")

                        load_dashboard_data.clear()
                        st.cache_data.clear()
                        record_upload(
                            admin_name,
                            target_name,
                            uploaded_file.name,
                            upload_action,
                            len(uploaded_df),
                        )
                        st.session_state.pending_dashboard = target_name
                        st.success(
                            f"{target_name} was {upload_action.lower()} successfully."
                        )
                        st.rerun()

        if admin_role in {"admin", "super_admin"}:
            if st.button("⚙️ Admin settings", use_container_width=True):
                show_admin_settings(admin_name, admin_role, all_dashboard_names)

# ------------------------------------------------------------------------------
# Data Loading & Processing
# ------------------------------------------------------------------------------
gallery_items = []

if selected_dashboard == OVERALL_DASHBOARD_NAME:
    combined_frames = []
    try:
        # Load all records from Supabase Database
        db_df = load_dashboard_data()
        if not db_df.empty:
            cleaned_db = clean_columns(db_df)
            combined_frames.append(cleaned_db)

        # Append any in-session uploaded overrides
        for name, item in st.session_state.uploaded_dashboards.items():
            uploaded_raw_df = load_data_file(
                BytesIO(item["data"]), item["filename"]
            )
            uploaded_df = add_source_label(uploaded_raw_df, name)
            combined_frames.append(uploaded_df)
            uploaded_gallery = extract_gallery_items(
                item["data"], item["filename"], uploaded_raw_df
            )
            item["gallery"] = uploaded_gallery
            gallery_items.extend(
                {
                    **photo,
                    "caption": f"{name} · {photo['caption']}",
                    "details": {
                        "Source Dashboard": name,
                        **photo.get("details", {}),
                    },
                }
                for photo in uploaded_gallery
            )
    except Exception as error:
        st.error(f"The overall dashboard data could not be loaded: {error}")
        st.stop()

    if combined_frames:
        df = pd.concat(combined_frames, ignore_index=True, sort=False)
    else:
        df = pd.DataFrame()
    source_label = f"{len(combined_frames):,} combined data source(s)"

else:
    if selected_dashboard in st.session_state.uploaded_dashboards:
        selected_item = st.session_state.uploaded_dashboards[selected_dashboard]
        source = BytesIO(selected_item["data"])
        source_label = selected_item["filename"]
        try:
            df = load_data_file(source, source_label)
            gallery_items = extract_gallery_items(
                selected_item["data"], selected_item["filename"], df
            )
            selected_item["gallery"] = gallery_items
        except Exception as error:
            st.error(f"The data file could not be loaded: {error}")
            st.stop()
    else:
        # After a browser refresh, restore the complete original CSV/XLSX from
        # Supabase Storage. Database rows contain only dashboard fields and do
        # not preserve every source column or embedded Excel image.
        stored_bytes, stored_name = download_dashboard_file_from_supabase(
            selected_dashboard
        )
        if stored_bytes and stored_name:
            try:
                df = load_data_file(BytesIO(stored_bytes), stored_name)
                gallery_items = extract_gallery_items(stored_bytes, stored_name, df)
                source_label = f"{stored_name} ({len(df)} records)"
            except Exception as error:
                st.warning(
                    f"The stored source file could not be restored; using database records: {error}"
                )
                db_df = load_dashboard_data(selected_dashboard)
                df = clean_columns(db_df) if not db_df.empty else pd.DataFrame()
                source_label = f"Database ({len(df)} records)"
        else:
            # Fallback for dashboards created before Storage persistence existed.
            db_df = load_dashboard_data(selected_dashboard)
            if not db_df.empty:
                df = clean_columns(db_df)
                source_label = f"Database ({len(df)} records)"
            else:
                df = pd.DataFrame()
                source_label = "No data available"

with left:
    st.markdown(
        f"""
        <div class="page-heading">
            <span class="eyebrow">INCIDENT COMMAND CENTER</span>
            <h1>{html.escape(str(selected_dashboard))} Dashboard</h1>
            <p>Current data: {html.escape(str(source_label))} · {len(df):,} records</p>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if selected_dashboard == OVERALL_DASHBOARD_NAME and not df.empty and "Source Dashboard" in df:
        data_sources = ", ".join(
            sorted(df["Source Dashboard"].dropna().astype(str).unique().tolist())
        )
        st.info(f"Data sources included in this Overall View: {data_sources}")

if df.empty:
    st.warning("No records found in the database or uploaded files. Please upload a dataset to populate the dashboard.")
    st.stop()

# ------------------------------------------------------------------------------
# Dashboard Metrics & Charts
# ------------------------------------------------------------------------------
incident_column = first_available_column(
    df, ["Type of Incident", "Service", "Accident Type", "Spot name"]
)
secondary_column = first_available_column(
    df,
    ["Traffic", "Severity", "Detail"],
)
municipality_column = first_available_column(df, ["Municipality"])
date_column = first_available_column(
    df, ["Registered Date", "Report Time", "Updated date and time", "Date"]
)
if date_column and not pd.api.types.is_datetime64_any_dtype(df[date_column]):
    df[date_column] = parse_dashboard_dates(df[date_column])

overall_record_count = 0
try:
    db_all = load_dashboard_data()
    overall_record_count = len(db_all)
except Exception:
    pass

for name, uploaded_item in st.session_state.uploaded_dashboards.items():
    if db_all.empty or name not in db_all.get("dashboard_name", []).values:
        try:
            overall_record_count += len(
                load_data_file(
                    BytesIO(uploaded_item["data"]), uploaded_item["filename"]
                )
            )
        except Exception:
            pass

with left:
    summary_chart1, summary_chart2 = st.columns(2, gap="large")
    with summary_chart1:
        if incident_column and df[incident_column].notna().any():
            incident_counts = category_counts(df, incident_column)
            incident_figure = px.pie(
                incident_counts,
                names="Category",
                values="Incidents",
                hole=0.42,
                title=f"Incidents by {incident_column}",
            )
            incident_figure.update_traces(
                textposition="inside", textinfo="percent+label"
            )
            incident_figure.update_layout(
                height=360,
                margin=dict(l=10, r=10, t=55, b=10),
                legend_title_text="",
            )
            st.plotly_chart(
                apply_pie_chart_theme(incident_figure), use_container_width=True
            )
        else:
            st.info("No incident type data is available.")

    with summary_chart2:
        if "Barangay" in df and df["Barangay"].notna().any():
            barangay_counts = category_counts(df, "Barangay")
            barangay_figure = px.pie(
                barangay_counts,
                names="Category",
                values="Incidents",
                hole=0.42,
                title="Incidents by Barangay",
            )
            barangay_figure.update_traces(
                textposition="inside", textinfo="percent+label"
            )
            barangay_figure.update_layout(
                height=360,
                margin=dict(l=10, r=10, t=55, b=10),
                legend_title_text="",
            )
            st.plotly_chart(
                apply_pie_chart_theme(barangay_figure), use_container_width=True
            )
        else:
            st.info("No barangay data is available.")

    selected_total, combined_total = st.columns(2, gap="large")
    selected_total.metric(
        f"{selected_dashboard} total records",
        f"{len(df):,}",
    )
    combined_total.metric(
        "Overall records (all dashboards)",
        f"{overall_record_count:,}",
    )

with st.expander("Filters", expanded=True):
    filter_count = 4
    if "Source Dashboard" in df:
        filter_count += 1
    if municipality_column:
        filter_count += 1
    filter_columns = st.columns(filter_count)
    filter_index = 0
    source_values = []
    if "Source Dashboard" in df:
        source_values = filter_columns[filter_index].multiselect(
            "Source Dashboard", options_for(df, "Source Dashboard")
        )
        filter_index += 1
    barangay_values = filter_columns[filter_index].multiselect(
        "Barangay", options_for(df, "Barangay")
    )
    filter_index += 1
    municipality_values = []
    if municipality_column:
        municipality_values = filter_columns[filter_index].multiselect(
            "Municipality",
            options_for(df, municipality_column),
            placeholder="All municipalities",
        )
        filter_index += 1
    incident_values = filter_columns[filter_index].multiselect(
        incident_column or "Incident type",
        options_for(df, incident_column) if incident_column else [],
    )
    filter_index += 1
    secondary_values = filter_columns[filter_index].multiselect(
        secondary_column or "Incident details",
        options_for(df, secondary_column) if secondary_column else [],
    )
    filter_index += 1
    registrant_values = filter_columns[filter_index].multiselect(
        "Registrant", options_for(df, "Registrant")
    )

    month_values = []
    week_values = []
    if date_column and df[date_column].notna().any():
        date_filter1, date_filter2 = st.columns(2)
        available_months = sorted(
            df.loc[df[date_column].notna(), date_column]
            .dt.to_period("M")
            .unique()
            .tolist(),
            reverse=True,
        )
        month_labels = {
            month: month.strftime("%B %Y")
            for month in available_months
        }
        selected_month_labels = date_filter1.multiselect(
            "Month",
            [month_labels[month] for month in available_months],
            placeholder="All months",
        )
        month_values = [
            month
            for month, label in month_labels.items()
            if label in selected_month_labels
        ]

        iso_calendar = df.loc[df[date_column].notna(), date_column].dt.isocalendar()
        available_weeks = sorted(
            (iso_calendar["year"].astype(str) + "-W" + iso_calendar["week"].astype(str).str.zfill(2))
            .unique()
            .tolist(),
            reverse=True,
        )
        week_values = date_filter2.multiselect(
            "Week",
            available_weeks,
            placeholder="All weeks",
        )

filtered = df.copy()
if date_column and month_values:
    filtered = filtered[
        filtered[date_column].dt.to_period("M").isin(month_values)
    ]
if date_column and week_values:
    filtered_iso = filtered[date_column].dt.isocalendar()
    filtered_week_labels = (
        filtered_iso["year"].astype(str)
        + "-W"
        + filtered_iso["week"].astype(str).str.zfill(2)
    )
    filtered = filtered[filtered_week_labels.isin(week_values)]

for column, chosen in {
    "Source Dashboard": source_values,
    "Barangay": barangay_values,
    municipality_column: municipality_values,
    incident_column: incident_values,
    secondary_column: secondary_values,
    "Registrant": registrant_values,
}.items():
    if column and chosen and column in filtered:
        filtered = filtered[filtered[column].map(normalize_category).isin(chosen)]

metric1, metric2, metric3, metric4 = st.columns(4)
metric1.metric("Total incidents", f"{len(filtered):,}")
metric2.metric(
    "Barangays",
    f"{filtered['Barangay'].map(normalize_category).nunique():,}"
    if "Barangay" in filtered
    else "—",
)
metric3.metric(
    "Reporting officers",
    f"{filtered['Reporting Officer'].map(normalize_category).nunique():,}"
    if "Reporting Officer" in filtered
    else "—",
)
latest = df[date_column].max() if date_column and date_column in df else pd.NaT
metric4.metric(
    "Latest report",
    latest.strftime("%d %b %Y") if pd.notna(latest) else "—",
)

# ------------------------------------------------------------------------------
# Tabs (Overview, Timeline, Gallery, Map, Records)
# ------------------------------------------------------------------------------
overview_tab, timeline_tab, gallery_tab, map_tab, records_tab = st.tabs(
    ["Overview", "Timeline", "Gallery", "Incident map", "Records"]
)

with overview_tab:
    chart1, chart2 = st.columns(2)
    with chart1:
        if date_column and filtered[date_column].notna().any():
            daily = (
                filtered.dropna(subset=[date_column])
                .assign(Date=lambda x: x[date_column].dt.floor("D"))
                .groupby("Date", as_index=False).size()
            )
            figure = px.line(
                daily, x="Date", y="size", markers=True, title="Incidents over time"
            )
            figure.update_layout(
                height=460,
                yaxis_title="Incidents",
                xaxis_title=None,
                xaxis=dict(type="date", tickformat="%b %Y", dtick="M2"),
            )
            figure.update_traces(
                line=dict(color="#78c5ff", width=2),
                marker=dict(color="#78c5ff", size=7),
                hovertemplate="%{x|%d/%m/%Y}<br>Incidents: %{y:,}<extra></extra>",
            )
            st.plotly_chart(apply_chart_theme(figure), use_container_width=True)
        else:
            st.info("No valid registered dates are available.")
    with chart2:
        if "Barangay" in filtered and filtered["Barangay"].notna().any():
            counts = category_counts(filtered, "Barangay", limit=12).sort_values("Incidents")
            figure = px.bar(
                counts,
                x="Incidents",
                y="Category",
                orientation="h",
                title="Incidents per Barangay",
            )
            st.plotly_chart(apply_chart_theme(figure), use_container_width=True)

    chart3, chart4 = st.columns(2)
    with chart3:
        if incident_column and filtered[incident_column].notna().any():
            counts = category_counts(filtered, incident_column, limit=20)
            figure = px.pie(
                counts,
                names="Category",
                values="Incidents",
                hole=0.45,
                title=f"Incidents by {incident_column}",
            )
            figure.update_traces(textposition="inside", textinfo="percent+label")
            st.plotly_chart(apply_pie_chart_theme(figure), use_container_width=True)
        else:
            st.info("No incident category column is available for the pie chart.")
    with chart4:
        if secondary_column and filtered[secondary_column].notna().any():
            counts = category_counts(filtered, secondary_column, limit=12)
            counts["Short Category"] = counts["Category"].map(shorten_category)
            counts = (
                counts.groupby("Short Category", as_index=False)["Incidents"]
                .sum()
                .sort_values("Incidents")
            )
            secondary_figure = px.bar(
                counts,
                x="Incidents",
                y="Short Category",
                orientation="h",
                color="Short Category",
                title=f"Incidents by {secondary_column}",
                labels={"Short Category": secondary_column},
            )
            st.plotly_chart(
                apply_chart_theme(secondary_figure),
                use_container_width=True,
            )
        else:
            st.info("No secondary incident category is available for this chart.")

with timeline_tab:
    st.subheader("Incident timeline")
    if date_column and filtered[date_column].notna().any():
        timeline_data = filtered.dropna(subset=[date_column]).copy()
        timeline_start = timeline_data[date_column].min()
        timeline_end = timeline_data[date_column].max()
        timeline_days = max((timeline_end - timeline_start).days, 0)

        timeline_control1, timeline_control2 = st.columns([2, 1])
        granularity_choice = timeline_control1.selectbox(
            "Timeline grouping",
            ["Automatic", "Daily", "Weekly", "Monthly"],
            help="Automatic chooses a grouping based on the selected date range.",
        )
        show_cumulative = timeline_control2.checkbox("Show cumulative total")

        if granularity_choice == "Automatic":
            resolved_granularity = (
                "Daily"
                if timeline_days <= 45
                else "Weekly"
                if timeline_days <= 365
                else "Monthly"
            )
        else:
            resolved_granularity = granularity_choice

        if resolved_granularity == "Daily":
            timeline_data["Timeline date"] = timeline_data[date_column].dt.floor("D")
        elif resolved_granularity == "Weekly":
            timeline_data["Timeline date"] = (
                timeline_data[date_column].dt.to_period("W").dt.start_time
            )
        else:
            timeline_data["Timeline date"] = (
                timeline_data[date_column].dt.to_period("M").dt.start_time
            )

        timeline_counts = (
            timeline_data.groupby("Timeline date", as_index=False)
            .size()
            .rename(columns={"size": "Incidents"})
            .sort_values("Timeline date")
        )
        y_column = "Incidents"
        if show_cumulative:
            timeline_counts["Cumulative incidents"] = timeline_counts[
                "Incidents"
            ].cumsum()
            y_column = "Cumulative incidents"

        timeline_figure = px.area(
            timeline_counts,
            x="Timeline date",
            y=y_column,
            markers=True,
            title=(
                f"{resolved_granularity} timeline · "
                f"{timeline_start:%d %b %Y} to {timeline_end:%d %b %Y}"
            ),
        )
        timeline_figure.update_traces(
            line=dict(color="#ff5964", width=3),
            fillcolor="rgba(255, 89, 100, 0.20)",
            hovertemplate="%{x|%d/%m/%Y}<br>Incidents: %{y:,}<extra></extra>",
        )
        timeline_figure.update_layout(
            height=460,
            xaxis_title="Registered Date (day/month/year)",
            xaxis=dict(type="date", tickformat="%d/%m/%Y"),
            yaxis_title=y_column,
            hovermode="x unified",
        )
        st.plotly_chart(
            apply_chart_theme(timeline_figure), use_container_width=True
        )
        st.caption(
            f"Using {date_column} · {len(timeline_data):,} dated record(s) · "
            f"Grouped {resolved_granularity.lower()}."
        )
    else:
        st.info("No valid date column is available for the timeline.")

with gallery_tab:
    st.subheader("Incident photo list")
    if gallery_items:
        st.caption(f"{len(gallery_items):,} photo(s) from the current dashboard data.")
        for photo_number, photo in enumerate(gallery_items):
            photo_details = photo.get("details", {})
            with st.container(border=True):
                photo_column, description_column = st.columns([1, 2], gap="large")
                with photo_column:
                    try:
                        image_source = html.escape(
                            gallery_image_source(photo["image"]), quote=True
                        )
                        modal_id = f"incident-photo-modal-{photo_number + 1}"
                        st.markdown(
                            f"""
                            <a href="#{modal_id}" title="View original-size photo">
                                <img src="{image_source}"
                                     alt="Incident photo {photo_number + 1}"
                                     style="width:260px;max-width:100%;height:180px;
                                            object-fit:cover;border-radius:14px;
                                            border:1px solid #3b4250;
                                            box-shadow:0 8px 24px rgba(0,0,0,.25);
                                            cursor:zoom-in;display:block;">
                            </a>
                            <div id="{modal_id}" class="photo-modal">
                                <a href="#" class="photo-modal-close"
                                   aria-label="Close photo">&times;</a>
                                <img src="{image_source}"
                                     alt="Original incident photo {photo_number + 1}">
                            </div>
                            """,
                            unsafe_allow_html=True,
                        )
                        st.caption("Click the photo to enlarge it.")
                    except Exception:
                        st.warning(
                            f"Photo {photo_number + 1} could not be displayed. "
                            "Check its image link or Excel attachment."
                        )
                with description_column:
                    st.markdown(f"#### Photo {photo_number + 1}")
                    if photo_details.get("Source Dashboard"):
                        st.caption(
                            f"Source dashboard: {photo_details['Source Dashboard']}"
                        )
                    detail_labels = (
                        "No",
                        "SpotName",
                        "Name",
                        "Title",
                        "Date",
                        "Address",
                    )
                    detail_rows = "".join(
                        (
                            '<div class="photo-details-label">'
                            f"{html.escape(detail_label)}</div>"
                            '<div class="photo-details-value">'
                            f"{html.escape(str(photo_details.get(detail_label) or '—'))}"
                            "</div>"
                        )
                        for detail_label in detail_labels
                    )
                    st.markdown(
                        f'<div class="photo-details-grid">{detail_rows}</div>',
                        unsafe_allow_html=True,
                    )
    else:
        st.info(
            "No photos were found. For CSV files, add a Photo or Image URL "
            "column containing http(s) links or data:image values. For XLSX "
            "files, you can insert images into any worksheet, including a "
            "separate Photos sheet."
        )

with map_tab:
    if {"latitude", "longitude"}.issubset(filtered.columns):
        map_data = filtered.dropna(subset=["latitude", "longitude"]).copy()
        if not map_data.empty:
            map_color_column = None
            if incident_column and incident_column in map_data:
                map_color_column = "Normalized Incident Category"
                map_data[map_color_column] = map_data[incident_column].map(
                    normalize_category
                )
                if "Source Dashboard" in map_data:
                    map_data[map_color_column] = (
                        map_data["Source Dashboard"].map(normalize_category)
                        + " — "
                        + map_data[map_color_column]
                    )
            hover_fields = [
                c
                for c in ["Spot name", "Barangay", incident_column, "Registered Date"]
                if c and c in map_data
            ]
            figure = px.scatter_map(
                map_data,
                lat="latitude",
                lon="longitude",
                hover_name="Spot name" if "Spot name" in map_data else None,
                hover_data=hover_fields,
                color=map_color_column,
                zoom=8,
                height=600,
            )
            figure.update_layout(
                map=dict(style="open-street-map"),
                margin=dict(l=0, r=0, t=0, b=0),
            )
            st.plotly_chart(apply_chart_theme(figure), use_container_width=True)
        else:
            st.info("No valid coordinates are available for the current filters.")
    else:
        st.info("The CSV needs latitude and longitude columns to display the map.")

with records_tab:
    st.dataframe(filtered, use_container_width=True, hide_index=True)
    st.download_button(
        "Download current dashboard data",
        filtered.to_csv(index=False).encode("utf-8-sig"),
        file_name=f"{selected_dashboard.lower().replace(' ', '_')}_records.csv",
        mime="text/csv",
    )
